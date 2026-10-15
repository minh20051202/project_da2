#!/usr/bin/env python3
"""Build Week-3 submission documents and a runnable Lean packet bundle."""

from __future__ import annotations

import shutil
import sys
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "Nop_Bao_Cao_Tung_Tuan" / "Tuan_03"
LEAN_OUT = OUT / "Lean_Packet"
PACKET_REL = Path("formalization/packets/VAT_ART14_C2_EXPORT_DOCS/v0.1.0")
LEAN_PACKET = ROOT / PACKET_REL
FILES = {
    "architecture": OUT / "DA2_NguyenTuanMinh_Tuan3_KienTrucModule_CauTrucDuLieu.docx",
    "api": OUT / "DA2_NguyenTuanMinh_Tuan3_DacTaAPI_Pseudocode.docx",
    "evaluation": OUT / "DA2_NguyenTuanMinh_Tuan3_Metrics_Eval_LeanMau.docx",
}
ZIP_OUT = OUT / "DA2_NguyenTuanMinh_Tuan3_LeanPacket.zip"

NAVY = "16324F"
BLUE = "2166AC"
TEAL = "168AAD"
LIGHT_BLUE = "EAF2F8"
LIGHT_GREY = "F2F4F6"
MID_GREY = "D8DEE4"
WHITE = "FFFFFF"


def shade(cell, color):
    props = cell._tc.get_or_add_tcPr()
    node = props.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        props.append(node)
    node.set(qn("w:fill"), color)


def margins(cell, top=75, start=90, bottom=75, end=90):
    props = cell._tc.get_or_add_tcPr()
    mar = props.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        props.append(mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        item = mar.find(qn("w:" + edge))
        if item is None:
            item = OxmlElement("w:" + edge)
            mar.append(item)
        item.set(qn("w:w"), str(value))
        item.set(qn("w:type"), "dxa")


def repeat_header(row):
    props = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    props.append(el)


def page_field(p):
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Trang ")
    r.font.name = "Times New Roman"
    r.font.size = Pt(9)
    begin, instr, end = OxmlElement("w:fldChar"), OxmlElement("w:instrText"), OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end.set(qn("w:fldCharType"), "end")
    r._r.extend([begin, instr, end])


def new_document(short_title):
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin, sec.bottom_margin = Cm(1.9), Cm(1.8)
    sec.left_margin, sec.right_margin = Cm(2.2), Cm(2.0)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for name, size, color in (("Title", 19, NAVY), ("Subtitle", 12, BLUE), ("Heading 1", 14, NAVY), ("Heading 2", 12, BLUE), ("Heading 3", 11, TEAL)):
        style = doc.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(9)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.keep_with_next = True
    header = sec.header.paragraphs[0]
    header.text = "ĐỒ ÁN 2 – NGUYỄN TUẤN MINH – TUẦN 03"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.name, run.font.size = "Arial", Pt(8)
        run.font.color.rgb = RGBColor.from_string("6B7785")
    page_field(sec.footer.paragraphs[0])
    props = doc.core_properties
    props.author = "Nguyễn Tuấn Minh"
    props.title = short_title
    props.subject = "Đồ án 2 – sản phẩm tuần 03"
    props.keywords = "Agentic RAG; Legal IR; PostgreSQL; Lean 4; đánh giá"
    return doc


def add_cover(doc, title, subtitle):
    for text, size, bold, color in (
        ("ĐẠI HỌC BÁCH KHOA HÀ NỘI", 12, True, NAVY),
        ("ĐỒ ÁN 2 – MI3390", 12, True, BLUE),
    ):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.font.name, r.font.size, r.bold = "Arial", Pt(size), bold
        r.font.color.rgb = RGBColor.from_string(color)
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(34), Pt(16)
    p.add_run(title)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(24)
    r = p.add_run(subtitle)
    r.italic, r.font.size = True, Pt(12)
    data = [
        ("Sinh viên", "Nguyễn Tuấn Minh"),
        ("MSSV", "20237465"),
        ("Giảng viên hướng dẫn", "PGS.TS. Trần Ngọc Thăng"),
        ("Thời điểm báo cáo", "Tuần 03 – ngày 08/10/2026"),
        ("Đề tài", "Agentic RAG hỗ trợ kiểm tra và hoàn thiện giải trình thuế GTGT đối với hoạt động xuất khẩu, tích hợp Lean 4"),
    ]
    table = doc.add_table(rows=len(data), cols=2)
    table.style, table.alignment = "Table Grid", WD_TABLE_ALIGNMENT.CENTER
    for i, (key, value) in enumerate(data):
        table.cell(i, 0).text, table.cell(i, 1).text = key, value
        shade(table.cell(i, 0), LIGHT_BLUE)
        for cell in table.rows[i].cells:
            margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        table.cell(i, 0).paragraphs[0].runs[0].bold = True
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(16)
    p.add_run("Hà Nội, 08/10/2026").italic = True
    doc.add_page_break()
    doc.add_heading("NỘI DUNG", level=1)


def h(doc, text, level=1):
    doc.add_heading(text, level=level)


def para(doc, text, style=None):
    return doc.add_paragraph(text, style=style)


def bullets(doc, items):
    for text in items:
        p = doc.add_paragraph(text, style="List Bullet")
        p.paragraph_format.space_after = Pt(2)


def numbered(doc, items):
    for text in items:
        p = doc.add_paragraph(text, style="List Number")
        p.paragraph_format.space_after = Pt(3)


def code(doc, text):
    style = doc.styles["Normal"]
    for line in text.strip("\n").splitlines():
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.35)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.keep_together = True
        r = p.add_run(line or " ")
        r.font.name = "Consolas"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor.from_string(NAVY)
    return style


def note(doc, lead, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade(cell, "FFF3E8")
    margins(cell, 100, 130, 100, 130)
    p = cell.paragraphs[0]
    p.add_run(lead + ": ").bold = True
    p.add_run(text)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def grid(doc, headers, rows, widths=None, font_size=8.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style, table.alignment = "Table Grid", WD_TABLE_ALIGNMENT.CENTER
    repeat_header(table.rows[0])
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = text
        shade(cell, NAVY)
        margins(cell)
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.name, run.font.size = "Arial", Pt(font_size)
            run.font.color.rgb = RGBColor.from_string(WHITE)
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            margins(cells[i])
            if ridx % 2:
                shade(cells[i], "F8FAFC")
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(1)
                for run in p.runs:
                    run.font.size = Pt(font_size)
        if widths:
            for i, width in enumerate(widths):
                cells[i].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def contents(doc, sections):
    bullets(doc, sections)
    doc.add_page_break()


def add_status_box(doc):
    note(doc, "Trạng thái bằng chứng", "Các quyết định dưới đây là thiết kế kiến trúc. Lean mẫu và verifier được chạy riêng; điều đó không có nghĩa PostgreSQL, RAG, Agent runtime hoặc output gate đã được triển khai.")


def build_architecture():
    doc = new_document("Kiến trúc mô-đun và cấu trúc dữ liệu – Tuần 03")
    add_cover(doc, "KIẾN TRÚC MÔ-ĐUN VÀ CẤU TRÚC DỮ LIỆU", "Thiết kế modular monolith cho Agentic RAG và Lean 4")
    contents(doc, [
        "Tóm tắt quyết định", "1. Mục tiêu, phạm vi và ranh giới", "2. Kiến trúc offline/online",
        "3. Tám mô-đun và quyền sở hữu", "4. Lưu trữ và mô hình dữ liệu",
        "5. RAG, Legal IR, evidence và manifest", "6. Snapshot, proof identity và audit",
        "7. Quyền an toàn và lộ trình",
    ])
    h(doc, "TÓM TẮT QUYẾT ĐỊNH")
    grid(doc, ["Hạng mục", "Thiết kế chốt"], [
        ("Kiểu ứng dụng", "Modular monolith; 8 mô-đun logic, chưa tách microservice."),
        ("Pháp luật", "Nguồn có phiên bản → Legal IR → lowering được rà soát → Lean 4 → phê duyệt → registry."),
        ("Online", "Case/evidence → approved rule → proof obligation → Lean → explanation → output gate."),
        ("Lưu trữ", "Một PostgreSQL database; schema core và retrieval; pgvector cùng database."),
        ("Tài liệu gốc", "PDF, ảnh, OCR, Lean/proof artifacts ở filesystem hoặc object storage; DB giữ URI/hash/provenance."),
        ("Quy tắc online", "Chỉ rule `approved`; retrieval semantic không được vượt lọc identity/version."),
        ("Trạng thái mẫu", "Article 14(2) demo đã compile; semantic review/approval vẫn pending."),
    ], [4.1, 12.4])
    add_status_box(doc)

    h(doc, "1. MỤC TIÊU, PHẠM VI VÀ RANH GIỚI")
    para(doc, "Thiết kế hệ thống hỗ trợ kiểm tra hồ sơ thuế GTGT cho hoạt động xuất khẩu, truy xuất căn cứ pháp luật theo thời điểm và tạo bản nháp giải trình có thể truy vết. Agent hỗ trợ người duyệt; không tự ban hành quyết định thuế.")
    grid(doc, ["Trong phạm vi thiết kế", "Ranh giới cần giữ"], [
        ("Văn bản pháp luật Việt Nam có phiên bản; hồ sơ, evidence, citations và audit.", "Lean kiểm tra suy diễn từ predicate/giả thiết được mã hóa; không tự xác thực OCR, chứng từ hay bản dịch Legal IR."),
        ("Rule rõ được lowering sang Lean; rule mơ hồ hoặc discretionary nêu boundary.", "Source retrieval không chứng minh corpus đầy đủ. Retrieval miss không đủ để gán evidence `absent`."),
        ("MVP cho phép case synthetic và dữ liệu đã ẩn danh.", "Không tự dùng candidate rule; không tự kết luận toàn bộ quyền khấu trừ từ một proof hẹp."),
    ], [8.1, 8.4])

    h(doc, "2. KIẾN TRÚC OFFLINE/ONLINE")
    h(doc, "2.1. Formalization ngoại tuyến", 2)
    code(doc, "Nguồn luật có phiên bản\n  → chia thành rule units và source spans\n  → Agent tạo Rule Legal IR candidate\n  → kiểm tra schema + semantic review\n  → lowering/template hoặc Lean candidate\n  → test + Lean check\n  → approval\n  → versioned formal-law registry")
    para(doc, "Formalization mới chỉ mang trạng thái candidate. Reviewer kiểm tra điều kiện, ngoại lệ, phủ định, chủ thể, thời gian, dẫn chiếu và boundary; technical pass không tự phê duyệt ngữ nghĩa.")
    h(doc, "2.2. Xử lý hồ sơ trực tuyến", 2)
    code(doc, "Case và tài liệu\n  → evidence extraction/review\n  → chọn approved rule đúng phiên bản\n  → proof-obligation builder\n  → premise retrieval + deterministic tactics\n  → Lean feedback/repair có ngân sách\n  → explanation + deterministic output gate + audit")
    para(doc, "Thay evidence sau khi khóa snapshot phải tạo run/snapshot mới. Claim nhiều phần được trả kết quả theo phạm vi từng claim; proof điều kiện chứng từ không mở rộng thành kết luận toàn bộ quyền khấu trừ.")

    h(doc, "3. TÁM MÔ-ĐUN VÀ QUYỀN SỞ HỮU")
    grid(doc, ["Mô-đun", "Chức năng và dữ liệu ghi", "Ranh giới"], [
        ("Source & Document", "Nguồn, artifact bất biến, phiên bản, document ingest/OCR và source spans.", "Không đổi bản nguồn gốc; thay đổi tạo artifact mới."),
        ("Legal Retrieval", "Keyword/full-text/vector tìm legal chunks và citations.", "Online chủ yếu đọc source/chunk; lọc version/status trước semantic ranking."),
        ("Formalization & Registry", "Legal IR, Lean source/manifest, tests, candidate/review/approval lifecycle.", "Chỉ reviewer có thẩm quyền chuyển rule sang approved."),
        ("Case & Evidence", "Case, documents, normalized evidence, provenance, review state.", "Unknown mặc định an toàn; Agent không tự sửa reviewed evidence."),
        ("Agent Orchestrator", "Tạo run, lập kế hoạch, gọi tools, quản lý budget và dừng.", "Không trực tiếp sửa approved rules, evidence đã review, proof cũ hay final result."),
        ("Proof & Lean", "Proof obligation, bindings, chạy Lean, diagnostics và proof artifact.", "Không thay theorem/rule version hoặc thêm giả thiết ngoài mapping được phép."),
        ("Explanation & Output", "Structured result, explanation draft, citations, deterministic output gate.", "Không phát hành claim verified nếu thiếu proof/identity/citation tương ứng."),
        ("Audit & Observability", "Event append-only, metrics và tham chiếu artifact.", "Lưu hành động/căn cứ ngắn và refs, không lưu hidden chain-of-thought."),
    ], [3.1, 7.0, 6.4], 8)
    para(doc, "Tám mô-đun là ranh giới code và quyền dữ liệu trong một ứng dụng. Việc một bảng thuộc mô-đun nào không đòi hỏi triển khai một service riêng.")

    h(doc, "4. LƯU TRỮ VÀ MÔ HÌNH DỮ LIỆU")
    grid(doc, ["Lớp", "Lưu trữ", "Nội dung"], [
        ("DB quan hệ", "PostgreSQL `project_da2`, schema `core`", "Nguồn/rule/case/evidence/run/proof/result/audit metadata và foreign keys."),
        ("Chỉ mục truy xuất", "Cùng DB, schema `retrieval`, extension pgvector", "Legal chunks, full-text/vector index; dữ liệu nguồn vẫn là artifact gốc."),
        ("Artifact store", "Filesystem hoặc object storage", "PDF, ảnh, OCR, source Lean, proof, diagnostics và output lớn; DB giữ URI/hash."),
    ], [3.0, 5.4, 8.1])
    note(doc, "Cùng một database", "`core` và `retrieval` là schema cùng PostgreSQL; pgvector không phải database thứ hai. Chưa yêu cầu Qdrant, Neo4j, Kafka hay Kubernetes cho MVP.")
    h(doc, "4.1. Mười một bảng logic tối thiểu", 2)
    grid(doc, ["Bảng", "Trách nhiệm và khóa/thông tin chính"], [
        ("`core.artifact`", "ID, URI, SHA-256, loại, kích thước, metadata. Trỏ tới file bất biến; nội dung đổi thì artifact/version mới."),
        ("`core.legal_source_version`", "Source/document ID, version, cơ quan, ngày công bố, hiệu lực từ/đến, trạng thái, artifact/hash."),
        ("`core.rule_version`", "Rule ID/version, Legal IR JSONB, trạng thái/effectivity, source version, manifest pointer/hash/summary."),
        ("`core.case`", "Case ID, chủ thể/ngữ cảnh giao dịch, ngày nghiệp vụ/kỳ thuế, trạng thái."),
        ("`core.case_document`", "Case ID, document type, artifact, trang/vùng, upload/processing metadata."),
        ("`core.evidence`", "Case/document refs, evidence type, status, raw/normalized values, provenance, confidence, review và reason."),
        ("`core.agent_run`", "Case ID, lifecycle, input snapshot/hash, budget và metrics tổng hợp."),
        ("`core.audit_event`", "Run/sequence/type/component, input/output refs, decision, metrics, time; append-only qua application contract."),
        ("`core.proof_artifact`", "Goal/theorem hash, rule/manifest/snapshot identity, premise bindings, toolchain, Lean result/artifacts."),
        ("`core.final_result`", "Run, claims/status/scope, proof/evidence/citation refs, unresolved items, output artifact."),
        ("`retrieval.legal_chunk`", "Source version, article/clause/point, exact text, embedding_text, locator/metadata, vector."),
    ], [4.3, 12.2], 8)
    h(doc, "4.2. Quan hệ cốt lõi", 2)
    code(doc, "legal_source_version 1─N rule_version 1─N proof_artifact\ncase 1─N case_document 1─N evidence\ncase 1─N agent_run 1─N audit_event\nagent_run 1─N proof_artifact 1─N final_result\nlegal_source_version 1─N legal_chunk")
    para(doc, "Không cần bảng riêng cho mọi candidate relation giữa claims, documents và evidence. Agent có thể suy luận liên kết tạm thời; links đã dùng thực sự cho proof/result được lưu trong binding artifact để tái lập. Relation table riêng chỉ bổ sung khi truy vấn/ownership đòi hỏi.")

    h(doc, "5. RAG, LEGAL IR, EVIDENCE VÀ MANIFEST")
    h(doc, "5.1. Chunking và retrieval", 2)
    bullets(doc, [
        "Chia theo Điều–Khoản–Điểm; điểm tự đủ nghĩa là ưu tiên. Không gộp chunks từ hai phiên bản luật.",
        "`text` giữ nguyên văn để trích dẫn. `embedding_text` được phép thêm tiêu đề/ngữ cảnh chỉ phục vụ embedding.",
        "MVP vectorize legal corpus. Evidence hồ sơ ưu tiên truy vấn có cấu trúc, không đẩy dữ liệu nhạy cảm vào legal vector index.",
        "Lọc source/version/effectivity trước; kết hợp full-text và vector, mở rộng dẫn chiếu, rerank. Similarity chỉ xếp hạng ứng viên, không chọn phiên bản thay registry.",
    ])
    h(doc, "5.2. Ba IR tách biệt", 2)
    grid(doc, ["IR", "Trả lời câu hỏi", "Thông tin tiêu biểu"], [
        ("Rule Legal IR", "Quy tắc nào phù hợp?", "Nguồn, thời gian, chủ thể, modality, điều kiện, kết luận, ngoại lệ, dẫn chiếu, boundary."),
        ("Case/Evidence IR", "Hồ sơ có dữ kiện gì?", "Evidence key/value/status, document/page provenance, review, conflict."),
        ("Proof Goal IR", "Mệnh đề Lean nào cần chứng minh?", "Target, typed hypotheses, premise bindings, identity và unresolved items."),
    ], [3.1, 4.0, 9.4])
    code(doc, '''{
  "rule_id": "VAT_ART14_C2_EXPORT_DOCS",
  "source": {"document_id": "114/VBHN-VPQH", "article": "14", "clause": "2", "point": null},
  "subject": "exported_goods_transaction",
  "conditions": [], "conclusion": "document_conditions_satisfied",
  "exceptions": [], "cross_references": [],
  "formalisation_boundary": [], "status": "candidate"
}''')
    para(doc, "Đây là hình dạng minh họa. Gói Legal IR của Article 14 hiện là JSON demo có trạng thái candidate; chưa phải schema/compiler tổng quát và chưa được reviewer phê duyệt.")
    h(doc, "5.3. Evidence và các trạng thái", 2)
    grid(doc, ["Trạng thái", "Ý nghĩa/điều kiện sử dụng"], [
        ("`present`", "Có evidence được ghi nhận cho slot; không tự khẳng định chứng từ xác thực hay đủ về pháp lý."),
        ("`absent`", "Thiếu có căn cứ/review; không được suy ra chỉ vì retrieval/OCR không tìm thấy."),
        ("`unknown`", "Chưa xác định; mặc định an toàn. Không gộp với absent hoặc notApplicable."),
        ("`notApplicable`", "Có căn cứ vì điều kiện cho phép miễn áp dụng; cần lưu reason/review ngoài Lean status slot."),
    ], [3.0, 13.5])
    para(doc, "Evidence giữ evidence ID, case/document ID, type, raw/normalized JSONB, provenance, confidence, review status và status reason. Confidence tách khỏi status. Khi có dữ kiện mâu thuẫn, trả các ứng viên để xử lý thay vì tự chọn giá trị có confidence cao nhất.")
    h(doc, "5.4. FormalizationManifest", 2)
    bullets(doc, [
        "JSON artifact đầy đủ, bất biến; chứa rule/version, module, predicate/theorems, imports, source/code/test hashes, toolchain và lineage.",
        "`rule_version` giữ manifest artifact pointer/hash và summary JSONB cho lookup phổ biến.",
        "Code, theorem hoặc field mapping đổi thì tạo manifest mới. Vector index có thể rebuild; manifest đã duyệt không sửa tại chỗ.",
    ])

    h(doc, "6. SNAPSHOT, PROOF IDENTITY VÀ AUDIT")
    bullets(doc, [
        "Mỗi `agent_run` khóa input snapshot/hash gồm thời gian nghiệp vụ, evidence IDs/status/value hashes và provenance hashes; MVP không cần bảng case_snapshot riêng.",
        "Proof khóa `theorem_statement_hash`, `rule_version_id`, `manifest_hash`, `case_snapshot_hash`, `premise_binding_hash` và toolchain.",
        "Evidence thay đổi → run/snapshot/proof/final result mới. Không ghi đè proof hay dùng proof cũ cho snapshot mới.",
        "Audit event theo run/sequence, giữ component, hành động, reason ngắn, refs và metrics. Nội dung lớn/nhạy cảm nằm ở artifact, audit giữ ID/hash; không cần blockchain/hash chain.",
        "Output gate kiểm tra identity, approved/effective version, no `sorry`/`admit`, axiom allowlist, citation và answer–proof agreement.",
    ])
    h(doc, "7. QUYỀN AN TOÀN VÀ LỘ TRÌNH")
    para(doc, "Agent là orchestrator, không có quyền ghi tùy ý vào registry. Reviewer phê duyệt rule; module Case quản lý evidence; module Proof phát hành proof artifact; Output module tạo final result; Audit append event. Các hợp đồng hiện là thiết kế, chưa có runtime/database implementation được xác nhận.")
    grid(doc, ["MVP bắt đầu", "Chỉ mở rộng nếu đo được nhu cầu"], [
        ("Một PostgreSQL + pgvector, filesystem/object store, 11 bảng logic, legal chunks và evidence có cấu trúc.", "Case vector index, graph database, microservices, queue/cluster và proof search plugin."),
        ("Approved rules, deterministic code/identity filters, bounded Agent–Lean loop, simple output templates.", "Thêm công nghệ hoặc model thứ hai theo ablation, năng lực và rủi ro đo được."),
    ], [8.1, 8.4])
    para(doc, "Cơ sở quyết định chi tiết và lịch sử thảo luận: `docs/architecture/ADR-001-legal-ir-formalization-and-proof-retrieval.md`.")
    return doc


def build_api_spec():
    doc = new_document("Đặc tả API và pseudocode – Tuần 03")
    add_cover(doc, "ĐẶC TẢ HỢP ĐỒNG API VÀ PSEUDOCODE", "Tool contracts cho luồng RAG, AI Agent và Lean 4")
    contents(doc, [
        "Tóm tắt và quy ước", "1. Hợp đồng tool nội bộ", "2. Luồng Orchestrator",
        "3. Vòng Agent–Lean", "4. Explanation và output gate", "5. Budget, lỗi và audit",
    ])
    h(doc, "TÓM TẮT VÀ QUY ƯỚC")
    para(doc, "MVP là modular monolith. Sáu hợp đồng dưới đây là hàm nội bộ; HTTP adapter có thể được bổ sung sau mà không đổi ý nghĩa dữ liệu. Đây là đặc tả thiết kế, chưa phải API đã triển khai.")
    add_status_box(doc)
    h(doc, "1. HỢP ĐỒNG TOOL NỘI BỘ")
    grid(doc, ["Tool", "Input chính", "Output/trạng thái"], [
        ("`search_legal_sources`", "run_id, query, source_refs?, max_results", "Chunks nguyên văn, locator, source/version/hash/cross-ref. `found`, `requires_context`, `no_match`."),
        ("`find_approved_rules`", "claim, case_context, legal_time_context, rule_refs?", "Legal IR, version, manifest, evidence needs, scope/boundary. `found`, `requires_context`, `no_match`."),
        ("`query_case_evidence`", "run_id, rule_version_id, evidence_keys?", "Per-slot status/value/provenance/review/reason; missing slot trả `unknown`."),
        ("`build_proof_obligation`", "run_id, rule_version_id, target_ref", "Goal/statement, bindings/imports/identity/unresolved. `ready`, `needs_evidence`, `requires_human`, `invalid_input`."),
        ("`retrieve_lean_premises`", "goal_id, proof_state_ref?, max_results", "Symbols, kind, type, code/source hash, dependencies, match basis. `found`, `no_match`, `invalid_context`."),
        ("`check_lean_proof`", "goal_id, proof_candidate_ref, timeout_ms", "Result/identity/diagnostics/artifact. `passed`, `proof_failed`, `rejected`, `timeout`, `runtime_error`."),
    ], [3.6, 5.6, 7.3], 8)
    h(doc, "1.1. Tìm legal source", 2)
    code(doc, '''search_legal_sources({
  "run_id": "run-123",
  "query": "Điều kiện chứng từ hàng xuất khẩu",
  "source_refs": [{"article": "14", "clause": "2"}],
  "max_results": 10
}) → {
  "status": "found",
  "chunks": [{
    "source_version_id": "...", "text": "...", "locator": "...",
    "source_ref": "...", "source_hash": "...",
    "temporal_metadata": {}, "cross_references": [], "retrieval_basis": []
  }]
}''')
    para(doc, "Lấy thời điểm nghiệp vụ từ run; thiếu mốc chọn phiên bản thì yêu cầu ngữ cảnh. Retrieval kết hợp keyword/full-text và vector sau lọc identity/version; citation chỉ dùng `text` nguyên văn, không dùng `embedding_text`. `no_match` chỉ nói corpus hiện tại chưa tìm thấy.")
    h(doc, "1.2. Tìm approved rule", 2)
    code(doc, '''find_approved_rules({
  "claim": "Hồ sơ đáp ứng điều kiện chứng từ cho hàng xuất khẩu",
  "case_context": {"transaction_type": "export_goods"},
  "legal_time_context": {"transaction_date": "<lấy từ case>"},
  "rule_refs": [{"article": "14", "clause": "2"}]
}) → {
  "status": "found",
  "candidates": [{"rule_version_id": "...", "manifest_ref": "...",
    "legal_ir": {}, "required_evidence": [], "scope": "...",
    "formalisation_boundary": [], "match_basis": []}],
  "uncovered_claims": []
}''')
    para(doc, "Nếu ngày nghiệp vụ thiếu, không tự đoán phiên bản. Rule match chưa có nghĩa rule chắc chắn áp dụng cho case; evidence và điều kiện áp dụng tiếp tục được kiểm tra. Nếu chỉ cover một phần claim, trả phần chưa cover.")
    h(doc, "1.3. Truy vấn evidence", 2)
    code(doc, '''query_case_evidence({
  "run_id": "run-123", "rule_version_id": "rule-v1",
  "evidence_keys": ["customsDeclaration"]
}) → {
  "slots": [{"requirement_key": "customsDeclaration", "evidence_id": null,
    "status": "unknown", "normalized_value": null, "provenance": [],
    "review_status": "pending", "status_reason": "Chưa xác định trong snapshot"}]
}''')
    para(doc, "Bỏ `evidence_keys` để trả mọi slot của rule. Evidence xung đột trả các ứng viên và reason; không auto-select. `present` không tự xác thực chứng từ. `unknown` không đổi thành `absent` chỉ vì tìm kiếm không thấy.")
    h(doc, "1.4. Dựng proof obligation", 2)
    code(doc, '''build_proof_obligation({
  "run_id": "run-123", "rule_version_id": "rule-v1",
  "target_ref": "VAT.Article14Clause2.SatisfiesDocumentConditions"
}) → {
  "status": "needs_evidence",
  "goal_id": "goal-456", "lean_statement": null,
  "identity": {"rule_version_id": "...", "manifest_hash": "...",
    "case_snapshot_hash": "...", "theorem_statement_hash": "..."},
  "unresolved_items": ["customsDeclaration:unknown"]
}''')
    para(doc, "Builder đọc field mapping từ manifest, kiểm tra snapshot và unresolved state rồi mới tạo goal. Không đưa chính kết luận cần chứng minh thành hypothesis. Positive goal cần evidence unknown thì `needs_evidence`; không bỏ điều kiện hoặc tự phủ định theorem.")
    h(doc, "1.5. Truy xuất Lean premise", 2)
    code(doc, '''retrieve_lean_premises({"goal_id": "goal-456",
  "proof_state_ref": "diagnostics-or-open-goals-ref", "max_results": 10
}) → {"status": "found", "candidates": [{
  "symbol": "VAT.Article14Clause2.insuranceCondition",
  "kind": "definition", "type_signature": "...",
  "code_ref": "formalization/Article14Clause2.lean", "source_hash": "...",
  "dependencies": [], "match_basis": ["goal/type match"]
}]}''')
    para(doc, "Lọc exact rule/version/type/import trước, mở rộng dependency, sau đó dùng similarity để bổ sung/rerank. Definition giúp unfold; theorem/lemma chỉ áp dụng nếu Lean xác nhận type. No-match không ngăn thử tactic cơ bản.")
    h(doc, "1.6. Kiểm tra proof", 2)
    code(doc, '''check_lean_proof({"goal_id": "goal-456",
  "proof_candidate_ref": "artifact://proof-candidate", "timeout_ms": 10000
}) → {"status": "proof_failed", "failure_kind": "unsolved_goals",
  "identity": {"theorem_statement_hash": "...", "case_snapshot_hash": "..."},
  "diagnostics_ref": "artifact://lean-diagnostics", "elapsed_ms": 320}''')
    para(doc, "Timeout ở ví dụ chỉ minh họa. `passed` đòi đúng theorem/identity, clean verification, proof hoàn chỉnh, không `sorry`/`admit`, axiom dependencies trong allowlist. Exit code 0 một mình chưa đủ. Open goals trả nếu runtime hỗ trợ; Lean CLI có thể khởi đầu bằng diagnostics.")

    h(doc, "2. LUỒNG ORCHESTRATOR")
    code(doc, '''run = create_run(case, transaction_time)
freeze_snapshot(run)
sources = search_legal_sources(run.id, claim, source_refs)
rules = find_approved_rules(claim, case_context, time_context, refs)
if rules.status == "requires_context": return request_context(rules)
if rules.status == "no_match": return unresolved("no_formalized_rule", sources)
rule = registry.validate_approved_version(agent_select(rules.candidates), run.time)
evidence = query_case_evidence(run.id, rule.version, rule.required_evidence)
goal = build_proof_obligation(run.id, rule.version, target_ref)
if goal.status != "ready": return gated_unresolved(goal, sources)
proof = proof_loop(goal, shared_budget)
decision = structured_claim_result(proof, evidence, sources)
draft = explain_from_records(decision)
return output_gate(draft, decision)''')
    note(doc, "Nhiều claim", "Tách claim và ghi scope/proof riêng; proof hẹp chỉ hỗ trợ claim tương ứng. Claim chưa cover vẫn ghi unresolved, không suy rộng.")

    h(doc, "3. VÒNG AGENT–LEAN CÓ NGÂN SÁCH")
    code(doc, '''candidate = deterministic_proof(fixed_goal)
while shared_budget.can_continue():
    result = check_lean_proof(fixed_goal.id, candidate,
                              shared_budget.allowed_timeout())
    audit.append(result)
    shared_budget.consume(result)
    if result.status == "passed": return proof_success(result)
    if result.status in {"rejected", "runtime_error"}: return unresolved(result)
    if not shared_budget.can_continue(): break
    premises = retrieve_lean_premises(fixed_goal.id, result.feedback_ref)
    shared_budget.consume(premises)
    if not shared_budget.can_continue(): break
    candidate = agent.propose_proof(
        fixed_goal=fixed_goal, feedback=result, premises=premises)
return unresolved("budget_exhausted")''')
    para(doc, "Budget chung tính lượt gọi/thời gian/chi phí theo cấu hình cho Agent, retriever và Lean; giới hạn từng lần Lean nằm dưới giới hạn run. Candidate repair không được đổi theorem, rule version hoặc snapshot. Evidence mới tạo run/snapshot mới, giữ lineage với run cũ.")

    h(doc, "4. EXPLANATION VÀ OUTPUT GATE")
    code(doc, '''{
  "claim_id": "claim-01", "target_ref": "...",
  "status": "verified_within_scope", "reason_code": "lean_passed",
  "scope": "document conditions in demo rule",
  "proof_refs": ["proof-123"], "evidence_refs": ["evidence-1"],
  "citations": [{"source_version_id": "...", "locator": "Article 14(2)(c)"}],
  "unresolved_items": ["other deduction requirements not checked"]
}''')
    bullets(doc, [
        "Claim verified phải có proof passed cho target tương ứng; rule/manifest/snapshot/proof identity phải khớp.",
        "Citation phải trỏ tới source version/locator hợp lệ và đoạn trích đúng; evidence refs phải khớp premise bindings.",
        "Giữ unresolved items và không nâng claim hẹp thành kết luận toàn bộ quyền khấu trừ.",
        "MVP dùng template cho kết luận chính; tự do diễn giải bổ sung cần được đánh giá về nguồn hỗ trợ.",
        "Proof fail/timeout không chứng minh phủ định goal. Output gate thất bại thì sửa draft có giới hạn, sau đó báo chưa xác minh/lỗi xử lý.",
    ])
    h(doc, "5. BUDGET, LỖI VÀ AUDIT")
    grid(doc, ["Tình huống", "Xử lý", "Trạng thái kết quả"], [
        ("Thiếu thời gian áp dụng", "Yêu cầu context; không đoán version.", "requires_context"),
        ("Chưa có approved formalization", "Ghi rõ phần chưa kiểm chứng bằng Lean; vẫn dẫn source nếu tìm thấy.", "unresolved/no_formalized_rule"),
        ("Evidence unknown/conflict", "Yêu cầu bổ sung hoặc người kiểm tra; snapshot mới khi evidence đổi.", "needs_evidence/requires_human"),
        ("Lean syntax/type/unsolved goals", "Premise retrieval hoặc sửa candidate; statement/identity giữ nguyên.", "proof_failed"),
        ("Timeout/budget exhausted", "Dừng vòng; lưu mức dùng và trạng thái chưa chứng minh.", "timeout/unresolved"),
        ("Identity/placeholder/axiom invalid", "Loại candidate; không xuất claim verified.", "rejected"),
        ("Runtime/gate error", "Lưu diagnostics; trả lỗi an toàn, không phát hành draft lỗi.", "runtime_error"),
    ], [4.2, 8.0, 4.5], 8.5)
    para(doc, "Audit ghi event/sequence/component, input/output artifact refs, action, reason ngắn, status, metrics và thời gian. Không ghi hidden chain-of-thought; prompt/raw output lớn lưu artifact phù hợp và hạn chế quyền.")
    return doc


def build_evaluation():
    doc = new_document("Metrics, evaluation protocol và Lean mẫu – Tuần 03")
    add_cover(doc, "METRICS, QUY TRÌNH ĐÁNH GIÁ VÀ LEAN MẪU", "Kế hoạch đánh giá A/B/C và packet formalization chạy được")
    contents(doc, [
        "Tóm tắt và trạng thái", "1. Hai bài toán đánh giá", "2. Bộ ca và nhãn vàng",
        "3. Metrics chính và chẩn đoán", "4. Cách đo formalization", "5. Quy trình và evaluation manifest",
        "6. Lean packet mẫu", "7. Ranh giới kết luận",
    ])
    h(doc, "TÓM TẮT VÀ TRẠNG THÁI")
    grid(doc, ["Phần", "Trạng thái"], [
        ("Kiến trúc/API/eval protocol", "Thiết kế trong ba tài liệu này."),
        ("Lean rule mẫu", "Lean source compile và sáu theorem type/axiom kiểm tra."),
        ("Source excerpts/mapping", "Khớp PDF theo trang; mapping là candidate cần semantic review."),
        ("Benchmark A/B/C, gold đầy đủ, 10–20 seeds", "Chưa thực hiện/chưa xác nhận hoàn tất."),
        ("Rule registry online", "Packet ở candidate, approval not_approved, online_eligible false."),
    ], [6.1, 10.4])
    note(doc, "Đọc kết quả đúng phạm vi", "`technical_checks_passed` chỉ mô tả các kiểm tra packet được chạy. Đây chưa phải semantic approval, active law registry hoặc kết luận pháp lý.")

    h(doc, "1. HAI BÀI TOÁN ĐÁNH GIÁ")
    h(doc, "1.1. Online case evaluation", 2)
    grid(doc, ["Cấu hình", "Thành phần", "Câu hỏi"], [
        ("A – RAG", "Retrieval một lượt + trả lời có citation/từ chối.", "Baseline RAG vận hành ra sao?"),
        ("B – Agentic RAG", "A + planning/tool calls/retrieval nhiều bước/đối chiếu evidence có giới hạn.", "Điều phối cải thiện phần thiếu/mâu thuẫn không?"),
        ("C – Agentic RAG + Lean", "B + approved rules, proof obligation, Lean, proof–answer gate.", "Lớp formalization/verification bổ sung cải thiện gì?"),
    ], [3.3, 8.4, 4.8])
    para(doc, "Dùng cùng corpus snapshot, case/question và model nền; khác biệt về tool/prompt/budget phải ghi rõ. Phép so sánh B→C đo cả lớp formalization, mapping, Lean và gates; chưa quy tác động riêng cho Lean kernel. Theo đặc tả tuần 2, cùng output evidence của B được ghép cặp cho nhánh Lean. Báo C trên toàn bộ tập và subset formalized được xác định từ gold/registry trước khi chạy.")
    h(doc, "1.2. Formalization evaluation riêng", 2)
    grid(doc, ["Variant", "Luồng", "Mục tiêu"], [
        ("F0", "Direct LLM → Lean", "Baseline sinh trực tiếp."),
        ("F1", "RAG examples → Lean", "Tác dụng retrieval seed."),
        ("F2", "RAG → Legal IR → Lean", "Tách lỗi semantic IR khỏi code."),
        ("F3", "RAG → Legal IR → template/domain lowering → Lean", "Đánh giá đường ưu tiên theo nhóm rule."),
    ], [2.4, 8.6, 5.5])
    para(doc, "Đơn vị ở đây là rule unit và held-out gồm điều khoản chưa làm seed/example. Held-out hồ sơ online và held-out formalization là hai bài test khác nhau.")

    h(doc, "2. BỘ CA VÀ NHÃN VÀNG")
    grid(doc, ["Split", "Mục tiêu tuần 2", "Cách dùng"], [
        ("Development", "40 case", "Prompt/schema/thresholds, phân tích lỗi và pilot."),
        ("Held-out", "20 case", "Đánh giá cuối; không dùng chỉnh cấu hình."),
        ("Lean-eligible", "Ít nhất 20 trong 60", "Phân bổ cả dev và held-out theo gold; không chọn theo kết quả C."),
        ("Reliability/security", "Ít nhất 15 tình huống riêng", "Prompt injection, version/permission, timeout/tool errors và recovery."),
    ], [3.3, 4.2, 9.0])
    para(doc, "Đây là quy mô mục tiêu trong đặc tả tuần 2; dataset hoàn thành chưa được xác nhận. Pilot mẫu quanh Article 14(2) gồm đủ evidence, absent, unknown, notApplicable có căn cứ, mâu thuẫn và ngoài scope/phiên bản. Biến thể của cùng case ở cùng split.")
    bullets(doc, [
        "Gold ghi source/version/locator/effective-time; evidence fields/status/provenance; missing/conflict.",
        "Ghi claims được phép và không được khẳng định, expected action (trả lời, hỏi thêm, chuyển người) và mức chắc chắn.",
        "Với Lean: expected target, premise mapping, obligation readiness/result và phản ví dụ nếu phù hợp.",
        "Hai người gắn nhãn độc lập và phân xử bất đồng; gold chưa review ghi draft, chưa gọi ground truth pháp lý.",
    ])

    h(doc, "3. METRICS CHÍNH VÀ CHẨN ĐOÁN")
    grid(doc, ["Nhóm", "Cách đo/đơn vị"], [
        ("Retrieval", "Recall@k trên căn cứ gold từng query; MRR/nDCG nếu cần thứ hạng. Source/rule đúng hiệu lực theo case."),
        ("Evidence", "Value extraction F1, status macro-F1; báo riêng nhầm unknown/absent/notApplicable; link correctness và conflict detection."),
        ("Citations", "Precision: citation hỗ trợ claim / citations được đánh giá; coverage: claim trọng yếu có citation hợp lệ / claims trọng yếu."),
        ("Conclusion", "Rubric 0–2 từng claim: đúng kết quả, scope và expected action. Bỏ claim vàng vẫn tính completeness; claim tự thêm cần căn cứ."),
        ("Sai khẳng định", "Đếm conclusion dứt khoát trái gold hoặc vượt scope; tách false acceptance kỹ thuật khỏi sai sót lọt qua final answer."),
        ("Abstention", "Selective accuracy đi cùng answer coverage; precision/recall hỏi thêm/chuyển người và false rejection."),
        ("Cost/robustness", "Latency p50/p95/max theo bước; tokens, model/tool calls, chi phí, retries, timeout, recovery và stability."),
    ], [3.3, 13.2], 8.5)
    h(doc, "3.1. Nhánh C", 2)
    bullets(doc, [
        "Rule coverage và proof readiness là hai mẫu số khác nhau. Unknown có thể thuộc rule scope nhưng chưa đủ evidence cho positive goal.",
        "Mapping correctness = bindings khớp evidence ID, value/status và field Lean theo gold.",
        "Clean-proof success = obligations gold xác định có thể chứng minh mà proof hoàn chỉnh/đúng identity vượt technical gate / toàn bộ obligations eligible đã khóa trước khi chạy. Timeout, builder error, exhausted budget giữ trong mẫu số.",
        "Proof failed/timeout không thành nhãn pháp lý phủ định. Sai PASS mục tiêu 0; kết quả thực nghiệm chưa có.",
    ])

    h(doc, "4. CÁCH ĐO FORMALIZATION")
    grid(doc, ["Metric", "Numerator / denominator"], [
        ("Semantic pass", "Rule units đúng toàn bộ điều kiện/kết luận/ngoại lệ/scope/time/cross-reference/human boundary trong scope / tổng test units."),
        ("First-pass compile", "Units compile ở lần đầu / tổng test units."),
        ("Final clean build", "Units đạt build sạch + identity/technical gates / tổng test units."),
        ("Boundary precision/recall", "Boundary đúng/tổng boundary system đề xuất và gold yêu cầu; báo N/A nếu gold không có boundary."),
        ("Mutation detection", "Mutants lỗi hợp lệ được phát hiện / tổng mutants lỗi hợp lệ."),
        ("Accepted packet", "Packets đạt semantic + technical + required tests / tổng test units."),
        ("Effort", "Repair rounds, runtime, human-review time trên mỗi unit."),
    ], [4.2, 12.3], 8.2)
    para(doc, "Mutation suite nên ưu tiên lỗi ngữ nghĩa còn typecheck được: bỏ điều kiện, AND→OR, chấp nhận unknown thay present, bỏ nhánh ngoại lệ. Compile thành công không chứng minh semantic fidelity.")

    h(doc, "5. QUY TRÌNH VÀ EVALUATION MANIFEST")
    numbered(doc, [
        "Tạo pilot/dev, kiểm tra schema/gold/scorer và chọn top-k, prompt, timeout, giới hạn bước trên development.",
        "Khóa manifest trước held-out: hashes/refs của dataset, gold, corpus, rule registry, code, prompt, model, retriever, toolchain và budget.",
        "Chạy A/B/C trên cùng đầu vào/thành phần nền; chạy ghép cặp theo case và giữ logs/raw outputs/artifact IDs.",
        "Gold chỉ dành cho scorer/review. Chấm tự động trường xác định; hai reviewer ẩn tên config chấm nghiệp vụ/ngôn ngữ và phân xử bất đồng.",
        "Báo counts/rates, subgroup, chất lượng–chi phí, uncertainty phù hợp. Variants cùng hồ sơ là một group; không giả định các biến thể độc lập.",
        "Timeout/tool errors giữ lại. Rerun hạ tầng theo policy định trước và giữ lần lỗi. Không tune held-out rồi dùng lại làm bằng chứng cải thiện.",
        "Phân tích lỗi theo source, retrieval, extraction, orchestration, formalization, proof, gate và explanation. Chưa chạy benchmark thì chỉ báo protocol/pilot.",
    ])

    h(doc, "6. LEAN PACKET MẪU")
    para(doc, "Packet: `Lean_Packet/formalization/packets/VAT_ART14_C2_EXPORT_DOCS/v0.1.0/`. Canonical artifacts và README cũng có ở `formalization/packets/VAT_ART14_C2_EXPORT_DOCS/v0.1.0/`; ZIP đi kèm tài liệu chứa self-contained copy và source PDF.")
    grid(doc, ["Thành phần", "Vai trò/trạng thái"], [
        ("source.json", "PDF SHA-256, điều/khoản, page/locators, trích đoạn điểm a/b/c, cross-refs; hiệu lực pending."),
        ("legal_ir.json", "Rule candidate gồm conditions, scope, exclusions và formalisation boundary."),
        ("mapping.json", "Chín field bindings; sáu synthetic cases/theorems."),
        ("manifest.json", "Version, namespace/symbols, artifact hashes, Lean v4.33.1, axiom allowlist."),
        ("Article14Clause2.lean", "Predicate kiểm tra trạng thái tài liệu cho scope demo; Lean 4 thuần, không import LegalLean."),
        ("verification.json", "Record command/exit code/hash/dependencies/mutation smoke results; semantic review pending."),
    ], [4.4, 12.1], 8.4)
    h(doc, "6.1. Scope và mapping mẫu", 2)
    para(doc, "Điểm a/b cùng điểm c cho hàng hóa xuất khẩu là scope demo đã mô hình hóa. Mỗi field dưới đây ánh xạ tới điều kiện trong predicate; điểm a gộp các lựa chọn chứng từ thành slot duy nhất nên loại chứng từ cần được xác nhận bên ngoài.")
    grid(doc, ["Lean field/evidence key", "Nguồn", "Điều kiện trong demo"], [
        ("inputVatInvoiceOrTaxReceipt", "14(2)(a)", "present"),
        ("nonCashPaymentForPurchase", "14(2)(b)", "present — ordinary branch"),
        ("foreignPartyContract", "14(2)(c)", "present"),
        ("salesInvoice", "14(2)(c)", "present"),
        ("nonCashPaymentForExport", "14(2)(c)", "present"),
        ("customsDeclaration", "14(2)(c)", "present"),
        ("packingList", "14(2)(c)", "present"),
        ("billOfLading", "14(2)(c)", "present"),
        ("insuranceCertificate", "14(2)(c)", "present OR notApplicable; applicability reason needs review"),
    ], [6.0, 2.5, 8.0], 8.2)
    note(doc, "Ranh giới", "Không bao gồm ngoại lệ thanh toán do Chính phủ quy định, xuất khẩu qua sàn/ca đặc thù, khoản 1 và khoản 3 Điều 14, quy định hướng dẫn, dịch vụ xuất khẩu, chứng thực chứng từ, ngày hiệu lực. Predicate true chỉ xác nhận điều kiện đã mã hóa trong demo.")
    h(doc, "6.2. Sáu synthetic tests và cách chạy", 2)
    grid(doc, ["ID", "Test", "Kỳ vọng"], [
        ("T01", "Mandatory slots present; insurance notApplicable", "Predicate true"),
        ("T02", "Customs unknown", "Predicate false; business layer asks for evidence"),
        ("T03", "Customs absent", "Predicate false in model"),
        ("T04", "Insurance unknown", "Predicate false"),
        ("T05", "Customs notApplicable", "Predicate false"),
        ("T06", "Insurance present; others complete", "Predicate true"),
    ], [1.4, 8.2, 6.9], 8.3)
    code(doc, '''# From the project root, or from the included Lean_Packet folder:
python3 scripts/verify_article14_packet.py

# Direct Lean compilation:
elan run leanprover/lean4:v4.33.1 lean formalization/Article14Clause2.lean''')
    para(doc, "Verifier read-only: validates artifact hashes, source excerpts against PDF pages 15–16 (printed pages 122–123), IR/mapping metadata, structure fields, Lean compile, theorem expected types and axiom dependencies. Three type-correct mutations—remove customs condition, permit unknown insurance, permit unknown mandatory evidence—are all caught by regression tests. This is a smoke check, not held-out benchmark.")
    h(doc, "7. RANH GIỚI KẾT LUẬN")
    bullets(doc, [
        "`technical_checks_passed`: các kiểm tra được ghi trong verification.json đã đạt trên đúng artifact hashes/toolchain.",
        "Semantic review/temporal applicability còn pending; lifecycle candidate, approval not_approved, online_eligible false.",
        "Legal IR/mapping JSON là demo, chưa phải schema/lowering compiler tổng quát. Module PostgreSQL, RAG, Agent loop và output gate vẫn là thiết kế.",
        "10–20 seed và các số liệu benchmark là kế hoạch, không phải kết quả đã đạt. Lean logic không xác minh tính trung thực của bản dịch hay tính xác thực chứng từ.",
    ])
    para(doc, "Đối chiếu kiến trúc và quyết định chi tiết trong `docs/architecture/ADR-001-legal-ir-formalization-and-proof-retrieval.md`.")
    return doc


def digest(path):
    import hashlib
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def bundle_packet():
    sources = (
        Path("resources/references/01-vbhn-vpqh-114.pdf"),
        Path("formalization/Article14Clause2.lean"),
        PACKET_REL,
        Path("scripts/verify_article14_packet.py"),
    )
    LEAN_OUT.mkdir(parents=True, exist_ok=True)
    for relative in sources:
        source = ROOT / relative
        if source.is_file():
            copies = [(source, LEAN_OUT / relative)]
        else:
            copies = [(path, LEAN_OUT / relative / path.relative_to(source)) for path in source.rglob("*") if path.is_file()]
        for original, destination in copies:
            destination.parent.mkdir(parents=True, exist_ok=True)
            if destination.exists():
                if digest(destination) != digest(original):
                    raise FileExistsError(f"Preserving existing, changed bundle file: {destination}")
            else:
                shutil.copy2(original, destination)
    with tempfile.TemporaryDirectory(prefix="week3-packet-") as temporary:
        temp_zip = Path(temporary) / ZIP_OUT.name
        with zipfile.ZipFile(temp_zip, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as archive:
            for path in sorted(LEAN_OUT.rglob("*")):
                if path.is_file():
                    archive.write(path, Path("Lean_Packet") / path.relative_to(LEAN_OUT))
        if ZIP_OUT.exists():
            if digest(ZIP_OUT) != digest(temp_zip):
                raise FileExistsError(f"Preserving existing, changed archive: {ZIP_OUT}")
        else:
            shutil.move(str(temp_zip), ZIP_OUT)


def validate_docx(path, required_phrases):
    if not path.exists() or path.stat().st_size < 10_000:
        raise ValueError(f"Missing or unexpectedly small DOCX: {path}")
    with zipfile.ZipFile(path) as archive:
        if archive.testzip() is not None or "word/document.xml" not in archive.namelist():
            raise ValueError(f"Invalid DOCX container: {path}")
    document = Document(path)
    text = docx_text(document)
    for phrase in required_phrases:
        if phrase not in text:
            raise ValueError(f"Required content missing in {path.name}: {phrase}")
    if len(document.sections) != 1 or len(document.tables) < 4:
        raise ValueError(f"Unexpected DOCX structure: {path}")


def docx_text(document):
    text = "\n".join(p.text for p in document.paragraphs)
    for table in document.tables:
        text += "\n" + "\n".join(cell.text for row in table.rows for cell in row.cells)
    return text


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    bundle_packet()
    documents = {
        "architecture": build_architecture(),
        "api": build_api_spec(),
        "evaluation": build_evaluation(),
    }
    for key, document in documents.items():
        target = FILES[key]
        if target.exists():
            old = Document(target)
            if docx_text(old) != docx_text(document):
                raise FileExistsError(f"Preserving existing, changed report: {target}")
        else:
            document.save(target)
    validate_docx(FILES["architecture"], ["Tám mô-đun", "core.artifact", "retrieval.legal_chunk"])
    validate_docx(FILES["api"], ["search_legal_sources", "check_lean_proof", "budget_exhausted"])
    validate_docx(FILES["evaluation"], ["Recall@k", "Semantic pass", "technical_checks_passed"])
    with zipfile.ZipFile(ZIP_OUT) as archive:
        if archive.testzip() is not None:
            raise ValueError("Corrupt Lean packet ZIP")
        names = set(archive.namelist())
        for required in (
            "Lean_Packet/resources/references/01-vbhn-vpqh-114.pdf",
            "Lean_Packet/scripts/verify_article14_packet.py",
            "Lean_Packet/formalization/packets/VAT_ART14_C2_EXPORT_DOCS/v0.1.0/README.md",
        ):
            if required not in names:
                raise ValueError(f"Missing bundle file: {required}")
    print("Created three Week-3 DOCX reports and the self-contained Lean packet.")
    for path in (*FILES.values(), ZIP_OUT):
        print(f"{path.relative_to(ROOT)}\t{path.stat().st_size} bytes")
    print(f"Lean_Packet files\t{sum(1 for path in LEAN_OUT.rglob('*') if path.is_file())}")


if __name__ == "__main__":
    sys.exit(main())

