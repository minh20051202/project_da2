#!/usr/bin/env python3
"""Create the single self-contained Week-4 upload folder. Never upload/submit it."""
from __future__ import annotations

import argparse
import json
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT = ROOT / "Nop_Bao_Cao_Tung_Tuan/Tuan 4"


def table(doc, heads, rows):
    t = doc.add_table(rows=1, cols=len(heads))
    t.style = "Table Grid"
    for cell, value in zip(t.rows[0].cells, heads):
        cell.text = str(value)
        cell.paragraphs[0].runs[0].bold = True
    for values in rows:
        for cell, value in zip(t.add_row().cells, values):
            cell.text = str(value)


def report(record, corpus, index):
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.left_margin, s.right_margin, s.top_margin, s.bottom_margin = Cm(2.2), Cm(2), Cm(1.9), Cm(1.9)
    doc.styles["Normal"].font.name = "Times New Roman"
    doc.styles["Normal"].font.size = Pt(11)
    doc.styles["Normal"].paragraph_format.space_after = Pt(5)
    doc.core_properties.author = "Nguyễn Tuấn Minh"
    doc.core_properties.title = "Báo cáo tiến độ Đồ án 2 — Tuần 04"
    doc.add_paragraph("ĐẠI HỌC BÁCH KHOA HÀ NỘI", style="Subtitle")
    doc.add_heading("BÁO CÁO TIẾN ĐỘ ĐỒ ÁN 2 — TUẦN 04", 0)
    doc.add_paragraph("Hạn nộp: trước 23h00 thứ Năm, 15/10/2026")
    table(doc, ["Thông tin", "Nội dung"], [
        ("Sinh viên / MSSV", "Nguyễn Tuấn Minh / 20237465"),
        ("GVHD", "PGS.TS. Trần Ngọc Thăng"),
        ("Đề tài", "Agentic RAG hỗ trợ kiểm tra và hoàn thiện giải trình thuế GTGT đối với hoạt động xuất khẩu, tích hợp Lean 4"),
        ("Phạm vi báo cáo", "Ba đầu ra tuần 04 ở mức bản nghiên cứu v0.1; không phải hệ thống tư vấn thuế production hoặc kết quả benchmark A/B/C."),
    ])
    doc.add_heading("1. Đối chiếu yêu cầu tuần 04", 1)
    table(doc, ["Đầu ra theo bảng tiến độ", "Nội dung bàn giao"], [
        ("Kho tài liệu pháp luật và nghiệp vụ v0.1", "37 PDF nguồn, bản chữ 3.424 trang, 2 tài liệu nghiệp vụ nội bộ; manifest/hash/phiên bản và cảnh báo chất lượng."),
        ("Mã nạp dữ liệu, chia đoạn, quản lý phiên bản, lập chỉ mục", f"Pipeline chạy được; {index['chunk_count']:,} chunk có locator; BM25, embedding E5 và hybrid RRF trên 5 nguồn/319 trang được chọn, không giả định đã index đủ 37 nguồn."),
        ("Bộ ca đánh giá đầu và hướng dẫn gắn nhãn", "16 biến thể giả lập thuộc một họ tình huống, tách đầu vào và nhãn dự thảo; hướng dẫn review và 11 fixture logic Lean offline."),
    ])
    doc.add_heading("2. Kho tài liệu và quản lý phiên bản", 1)
    doc.add_paragraph("Giữ nguyên PDF gốc; đối chiếu SHA-256 với kiểm kê nguồn trước đó và trích xuất bằng pdftotext -layout. Bản chữ theo trang PDF từ 1, giữ nguyên khoảng trắng/xuống dòng, hash từng trang. Tên và URL kế thừa danh mục tuần 01 có provenance; không coi metadata chưa kiểm chứng là thông tin hiệu lực đã duyệt.")
    doc.add_paragraph("Corpus source snapshot v0.1.2 chỉ gồm manifest.json, pages.jsonl và README.md. Manifest gộp danh mục, thống kê, metadata, cảnh báo/review; không tạo CSV hoặc JSON trùng lặp cho từng bảng. Thay đổi nguồn tạo version mới, không ghi đè snapshot. Hash kiểm tra tính toàn vẹn, không xác nhận hiệu lực hoặc tính xác thực pháp lý.")
    table(doc, ["Nhóm", "Giới hạn công khai"], [
        ("03, 07, 13", "Cần OCR; chưa xem bản chữ ít ỏi là toàn văn."),
        ("20, 29, 30, 31, 33, 36, 37", "Có chỉ dẫn phần tiếp ở số Công báo khác; tệp 33 thiếu phần tiếp phụ lục."),
        ("14, 19, 33", "PDF chứa nhiều văn bản; chưa tự đưa vào index khi ranh giới chưa review."),
        ("32", "Chỉ có TT 21/2026 sửa TT 80/2021, chưa có toàn văn TT 80/2021."),
        ("24–27", "Quản trị AI/dữ liệu, không dùng làm căn cứ trả lời nghiệp vụ VAT."),
        ("Tất cả nguồn", "Semantic/temporal review pending; ngày chưa xác minh để null; online_eligible=false."),
    ])
    doc.add_paragraph("Hai tài liệu nghiệp vụ là tài liệu thiết kế nội bộ, không phải hướng dẫn thuế chính thức hoặc hồ sơ doanh nghiệp. PDF nghiên cứu kỹ thuật không đưa vào corpus căn cứ pháp lý. Số trang/chunk/mốc cấu trúc không được gọi là tổng quy tắc pháp luật.")
    doc.add_heading("3. Pipeline chunking và retrieval chạy được", 1)
    for line in [
        "Nạp source snapshot đã khóa hash; chỉ lấy legal_pdf_page, không vector hóa tài liệu nghiệp vụ hoặc nhãn đánh giá.",
        "Chia theo Điều–Khoản–Điểm bằng parser cấu trúc ứng viên; giữ phần dẫn nhập cha. Đoạn quá dài chia thêm thành phần có cùng metadata cha; text trích dẫn luôn là span nguyên văn có char offsets và trang PDF.",
        "embedding_text bổ sung tiêu đề/ngữ cảnh chỉ phục vụ tìm kiếm. Không lấy embedding_text thay cho nguyên văn để citation.",
        "BM25 dùng tokenizer từ tiếng Việt bỏ dấu để tìm từ khóa; semantic dùng mô hình đa ngôn ngữ thật intfloat/multilingual-e5-small, ONNX int8 chính thức, prefix query:/passage:, attention-mask mean pooling và L2 normalization.",
        "Lọc source snapshot trước xếp hạng; hybrid kết hợp hai danh sách bằng RRF. Chế độ offline phải được yêu cầu rõ; online cần ngày nghiệp vụ và dữ liệu approved/đúng hiệu lực, hiện bị chặn với mọi source candidate.",
        "Chỉ mục file có thể kiểm tra/tái tạo độc lập; PostgreSQL/pgvector vẫn là kiến trúc triển khai dự kiến, chưa được triển khai production ở tuần 04. Không thay đổi ADR thành SQLite hoặc vector database khác.",
    ]:
        doc.add_paragraph(line, style="List Number")
    table(doc, ["Thông số thực tế", "Giá trị"], [
        ("Nguồn trong index", ", ".join(index["selected_source_versions"])),
        ("Số trang / chunk / chiều vector", f"{index['pdf_pages']} / {index['chunk_count']} / {index['dimension']}"),
        ("Nguồn model", "https://huggingface.co/intfloat/multilingual-e5-small"),
        ("Revision model", index["model"]["revision"]),
        ("Giới hạn embedding", f"512 token; {index['truncated_embedding_inputs']} đầu vào embedding có truncation. Nguyên văn citation không bị cắt bởi tokenizer."),
        ("Định danh index", index["index_id"]),
    ])
    doc.add_paragraph("Parser tự động chưa thay thế rà soát cấu trúc/ngữ nghĩa: footnote, đoạn sửa đổi, dẫn chiếu và ngoại lệ có thể cần review. Tất cả chunk là structural_candidate. Việc một source chưa có cảnh báo ưu tiên không chứng minh nguồn đầy đủ hoặc còn hiệu lực.")
    doc.add_heading("4. Bộ ca pilot và quy trình gắn nhãn", 1)
    doc.add_paragraph("16 ca là biến thể của một họ tình huống giả lập quanh Điều 14 khoản 2, không phải 16 hồ sơ độc lập hoặc toàn bộ 60 ca mục tiêu. Tất cả thuộc pilot_dev, không có held-out. Tình huống gồm đủ/thiếu/chưa rõ chứng từ, N/A có/không có lý do, mâu thuẫn, chỉ dẫn độc hại trong tài liệu, ngoài scope và thiếu ngữ cảnh thời gian.")
    table(doc, ["Trạng thái chứng cứ", "Cách gắn nhãn"], [
        ("present", "Có đoạn tài liệu/giả định kiểm thử phù hợp và locator. Không có nghĩa chứng từ thật."),
        ("absent", "Có xác nhận thiếu rõ ràng; không suy ra từ search/OCR không tìm thấy."),
        ("unknown", "Chưa rõ, còn mâu thuẫn hoặc N/A tự khai chưa đủ căn cứ."),
        ("notApplicable", "Có lý do và nhánh mô hình cho phép; trong pilot là giả định bảo hiểm không phát sinh."),
    ])
    doc.add_paragraph("cases.jsonl chỉ chứa đầu vào; labels.jsonl dành riêng cho scorer/reviewer, không đưa vào prompt/index. Các nhãn đều draft do assistant đề xuất; reviewer 1/2/adjudicator đều null. Cần hai người gắn nhãn độc lập và phân xử bất đồng trước khi gọi là ground truth nghiệp vụ. Biến thể cùng họ không được chia sang held-out.")
    doc.add_paragraph("11 ca có fixture Lean offline; 5 ca chặn bind được báo riêng. Proof true/false chỉ nói về predicate tài liệu đã mã hóa; không suy ra được quyền khấu trừ/hoàn thuế cuối cùng hoặc tính trung thành của bản mã hóa với luật.")
    doc.add_heading("5. Kiểm tra đã chạy và bằng chứng", 1)
    table(doc, ["Kiểm tra", "Kết quả đã quan sát"], [
        ("Unit tests", f"{record['unit_tests']['count']} kiểm thử qua; gồm toàn vẹn nguồn, provenance, chống lộ nhãn, chunking, bộ lọc và action/fixture gates."),
        ("Corpus", "42 artifact theo manifest và 3.424 trang/828 đoạn nội bộ qua kiểm tra hash/locator."),
        ("Chunk/citation", f"Tái sinh toàn bộ {index['chunk_count']} chunk và đối chiếu đúng frozen index; ma trận vector hợp lệ, chuẩn hóa L2."),
        ("Retrieval smoke", f"3 anchor đã biết × 3 cách retrieval = 9 lượt; anchor đúng trong top-5 ở {record['retrieval']['anchor_hits_top5']}/9 lượt. Đây không phải held-out Recall@5 hoặc benchmark 3 agent."),
        ("Lean offline", f"{record['lean'].get('checked_fixtures', 0)} theorem fixture và seed packet/mutation smoke checks qua; semantic/temporal approval vẫn pending."),
        ("Chạy độc lập", "Các kiểm tra chạy lại bên trong bản thư mục bàn giao tạm, không cần cache, source tree gốc hoặc tải model từ mạng."),
    ])
    doc.add_paragraph("Chi tiết lệnh, exit code, hash/ID, probe và axiom diagnostics: KiemTra/ket_qua.json. Kết quả chỉ xác nhận kiểm tra kỹ thuật đã nêu, không phải legal approval hoặc hiệu quả với kế toán. Chưa có thực nghiệm A/B/C, OCR toàn bộ corpus, nhãn chuyên gia, held-out hoặc ứng dụng tư vấn hoàn chỉnh.")
    doc.add_heading("6. Cách sử dụng bộ nộp và kế hoạch tiếp theo", 1)
    doc.add_paragraph("Mở README.md tại gốc thư mục Tuan 4. Đã kèm PDF, model ONNX/tokenizer, dữ liệu/index, Lean seed, mã nguồn, kiểm thử và requirements. Python 3.12 và các package cần được cài trước khi chạy; Lean/elan chỉ cần cho kiểm tra proof offline, Poppler chỉ cần khi tái trích xuất PDF.")
    for cmd in ["python -m pip install -r requirements-week4.txt", "python -B scripts/verify_week4_submission.py", "python -B scripts/verify_week4_submission.py --with-lean"]:
        doc.add_paragraph(cmd)
    doc.add_paragraph("Tuần tiếp theo: hoàn thiện baseline RAG sinh câu trả lời có trích dẫn/từ chối, phân tích lỗi retrieval, rà soát chunk và hiệu lực, mở rộng hồ sơ độc lập và hoàn thiện nhãn theo reviewer. Không khẳng định Lean tốt hơn baseline khi chưa chạy đối chứng.")
    doc.add_paragraph("Bộ này được chuẩn bị để sinh viên tải cả thư mục lên; chưa thực hiện nộp bài hoặc đổi trạng thái trong bảng tiến độ.")
    return doc


def readme(record, index):
    return f"""# Đồ án 2 — bộ nộp tuần 04

Sinh viên: Nguyễn Tuấn Minh, MSSV 20237465. GVHD: PGS.TS. Trần Ngọc Thăng.
Hạn nộp theo bảng tiến độ: trước 23h00 thứ Năm 15/10/2026.
**Tải nguyên thư mục `Tuan 4` lên, không cần ZIP.** Chưa tự nộp hoặc đánh dấu đã nộp.

## Đủ ba nhóm đầu ra tuần 04 ở mức v0.1 nghiên cứu

1. `resources/` + `data/corpus/v0.1.2/`: 37 PDF/3.424 trang, tài liệu nghiệp vụ
   nội bộ, manifest/hash, bản chữ và cảnh báo chất lượng.
2. `scripts/` + `data/retrieval/v0.1/`: mã nạp/chia đoạn/quản lý version;
   {index['chunk_count']} chunk trên 5 nguồn/319 trang; BM25, E5 semantic và hybrid RRF.
3. `data/evaluation/pilot-v0.1/`: 16 biến thể giả lập, input/nhãn dự thảo tách
   riêng, hướng dẫn gắn nhãn và 11 fixture Lean offline. Không có held-out.

`DA2_NguyenTuanMinh_Tuan4_BaoCaoTienDo.docx` là báo cáo tổng hợp duy nhất.
`KiemTra/ket_qua.json` gộp toàn bộ kết quả kiểm tra, không sinh log riêng từng ca.
`models/` có ONNX và tokenizer đã tải, nên tìm kiếm chạy offline sau khi cài package.
`formalization/` giữ seed candidate và provenance hiện có, không tạo/phê duyệt luật mới.
`Nop_Bao_Cao_Tung_Tuan/Tuan_01` và `Tuan_02` bên trong bộ này chỉ chứa **3 DOCX
tham chiếu cần thiết** theo đường dẫn manifest, không phải bản sao toàn bộ tuần cũ.

## Cài và chạy từ chính thư mục này

Yêu cầu đã kiểm tra: Python 3.12, CPU, các package trong requirements; không
cần API key, PostgreSQL hoặc kết nối mạng khi truy xuất bằng model đã kèm.
Lần cài package cần mạng nếu máy chưa có package. Không kèm virtualenv/runtimes
riêng của máy phát triển vì không dùng chung được giữa các hệ điều hành.

```bash
python -m venv .venv
# Windows PowerShell: .venv\\Scripts\\Activate.ps1
# Linux/macOS: source .venv/bin/activate
python -m pip install -r requirements-week4.txt
python -B scripts/verify_week4_submission.py
```

Nếu có elan/Lean 4.33.1, kiểm tra thêm 11 fixture và seed packet:

```bash
python -B scripts/verify_week4_submission.py --with-lean
```

Nếu chưa có toolchain, lệnh không có `--with-lean` vẫn kiểm tra corpus, pilot,
index, unit tests và retrieval; trường lean ghi rõ chưa chạy, không giả báo pass.
Poppler/pdftotext chỉ cần khi tái trích xuất corpus hoặc kiểm tra seed packet đầy đủ.

Tìm kiếm thử, không sinh câu trả lời pháp lý:

```bash
python -B scripts/legal_retrieval.py search --offline-snapshot --source-id LEGAL-001 --mode keyword --query "chứng từ thanh toán không dùng tiền mặt"
python -B scripts/legal_retrieval.py search --offline-snapshot --source-id LEGAL-001 --mode semantic --query "xuất khẩu cần những giấy tờ nào"
python -B scripts/legal_retrieval.py search --offline-snapshot --source-id LEGAL-001 --mode hybrid --query "tờ khai hải quan phiếu đóng gói vận đơn"
```

Các source IDs trong index: {', '.join(index['selected_source_versions'])}.
Citation trả về text nguyên văn, article/clause/point, source version/hash và
trang PDF. embedding_text chỉ để retrieval, không phải câu trích dẫn.

Kiểm tra/tái tạo dữ liệu (lệnh build chỉ chạy khi muốn tái tạo, tốn thời gian):

```bash
python -B scripts/build_corpus_release.py --verify-only
python -B scripts/build_corpus_release.py
python -B scripts/legal_retrieval.py verify
python -B scripts/legal_retrieval.py build
```

Script từ chối ghi đè bản đóng băng khác nội dung. Khi dữ liệu/công cụ làm thay
đổi artifact, cần review và tạo version mới; không xóa version cũ để ép pass.

## Kết quả thực tế đã lưu

- {record['unit_tests']['count']} unit tests đạt.
- Regenerate/đối chiếu {index['chunk_count']} chunk, hash nguồn và model; vector {index['dimension']} chiều chuẩn L2.
- 3 anchor đã biết × 3 chế độ = 9 smoke searches; đúng anchor trong top-5 ở
  {record['retrieval']['anchor_hits_top5']}/9 lượt. Không gọi đây là held-out accuracy/Recall@5.
- 11 theorem fixture offline và seed packet checks qua; không phải semantic approval.

## Giới hạn phải đọc trước khi dùng

- Đây là bản v0.1 nghiên cứu, không phải hệ thống tư vấn thuế production.
- Không tự index toàn bộ 37 nguồn: scan 03/07/13 chưa OCR; có nguồn thiếu phần
  tiếp hoặc chứa nhiều văn bản. Cảnh báo và danh sách không index ở manifest.
- Parser Điều–Khoản–Điểm là candidate, cần rà footnote/sửa đổi/dẫn chiếu; vector
  có tối đa 512 token đầu vào, {index['truncated_embedding_inputs']} embedding input có truncation. Text citation vẫn nguyên vẹn.
- Metadata hiệu lực chưa duyệt để null. Chế độ online bị chặn; chỉ chạy snapshot
  khi có cờ offline rõ ràng. Không suy ra quyền hoàn thuế từ proof tài liệu.
- PostgreSQL/pgvector vẫn là thiết kế deployment; file index này là artifact
  offline tái tạo được, không tuyên bố đã triển khai database đó hoặc đổi ADR.
- Pilot là một họ tình huống, mọi nhãn draft; không có 16 hồ sơ độc lập, expert
  gold, held-out, OCR benchmark hoặc kết quả A/B/C. Agent chỉ đọc cases.jsonl,
  không đọc labels.jsonl hoặc guide chứa đáp án.
- README trong snapshot corpus/pilot giữ mốc lịch sử lúc tạo, có thể nhắc đường
  dẫn bộ bàn giao cũ; **đường dẫn/cách chạy hiện tại theo README gốc này**.

## Nguồn kỹ thuật và bản quyền model

Model E5 chính thức: https://huggingface.co/intfloat/multilingual-e5-small
(license MIT theo model card đi kèm), revision `{index['model']['revision']}`.
Giữ nguyên MODEL_CARD.md và hash các asset trong manifest index. Không gửi hồ
sơ lên dịch vụ embedding bên ngoài. `scripts/download_embedding_model.py` chỉ
dùng khi cần tải lại đúng revision đã khóa; model đã có trong thư mục này.
""".encode()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    target = args.output.resolve()
    if target.exists():
        raise ValueError("Output already exists; preserve the old submission before replacing it")
    target.parent.mkdir(parents=True, exist_ok=True)
    corpus = json.loads((ROOT / "data/corpus/v0.1.2/manifest.json").read_text())
    paths = {ROOT / item["path"] for item in corpus["original_artifacts"]}
    for rel in ("data/corpus/v0.1.2", "data/evaluation/pilot-v0.1", "data/retrieval/v0.1", "models/multilingual-e5-small", "formalization"):
        paths.update(p for p in (ROOT / rel).rglob("*") if p.is_file() and "__pycache__" not in p.parts)
    for name in ("build_corpus_release.py", "download_embedding_model.py", "legal_retrieval.py", "validate_pilot_cases.py",
                 "verify_article14_packet.py", "verify_week4_submission.py", "build_week4_submission.py"):
        paths.add(ROOT / "scripts" / name)
    paths.update((ROOT / "tests").glob("test_*.py"))
    paths.add(ROOT / "requirements-week4.txt")
    with tempfile.TemporaryDirectory(prefix="da2-week4-final-", dir=target.parent) as temp:
        stage = Path(temp) / "Tuan 4"
        stage.mkdir()
        for source in sorted(paths):
            destination = stage / source.relative_to(ROOT)
            destination.parent.mkdir(parents=True, exist_ok=True)
            shutil.copyfile(source, destination)
        cmd = [sys.executable, "-B", "scripts/verify_week4_submission.py", "--with-lean", "--output-json", "KiemTra/ket_qua.json"]
        checks = subprocess.run(cmd, cwd=stage, capture_output=True, text=True, timeout=180)
        if checks.returncode:
            raise RuntimeError("Clean-folder verification failed: " + checks.stdout + checks.stderr)
        record = json.loads((stage / "KiemTra/ket_qua.json").read_text())
        index = json.loads((stage / "data/retrieval/v0.1/manifest.json").read_text())
        (stage / "README.md").write_bytes(readme(record, index))
        doc = report(record, corpus, index)
        doc_path = stage / "DA2_NguyenTuanMinh_Tuan4_BaoCaoTienDo.docx"
        doc.save(doc_path)
        with zipfile.ZipFile(doc_path) as archive:
            if archive.testzip():
                raise ValueError("Corrupt DOCX")
        reopened = Document(doc_path)
        text = "\n".join(p.text for p in reopened.paragraphs)
        if "ĐẠI HỌC BÁCH KHOA HÀ NỘI" not in text or "TUẦN 04" not in text:
            raise ValueError("Wrong report identity")
        if list(stage.rglob("*.zip")) or list(stage.rglob("*.pyc")):
            raise ValueError("Unexpected archive/cache in the submission")
        stage.rename(target)
    print(json.dumps({"submission_folder": str(target), "status": record["status"],
                      "unit_tests": record["unit_tests"]["count"], "retrieval_chunks": index["chunk_count"],
                      "smoke_anchor_hits": record["retrieval"]["anchor_hits_top5"],
                      "lean_fixtures": record["lean"]["checked_fixtures"], "submitted": False}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
